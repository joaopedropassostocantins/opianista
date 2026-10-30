#!/usr/bin/env python3
"""Gera o relatorio preliminar em .docx conforme ABNT (NBR 14724, 6024, 6027, 6028).

Le o Markdown do relatorio (titulo '#', secoes '##', subsecoes '###',
paragrafos, listas numeradas, tabelas '|') e produz: capa, folha de rosto,
resumo, sumario com paginas medidas numa primeira renderizacao (LibreOffice), texto com paginacao no
canto superior direito a partir da introducao, e referencias.
Requer python-docx.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_TAB_LEADER
from docx.shared import Cm, Pt, RGBColor
import subprocess
import tempfile

FONT = "Arial"


def set_margins(section):
    section.top_margin, section.left_margin = Cm(3), Cm(3)
    section.bottom_margin, section.right_margin = Cm(2), Cm(2)
    section.page_height, section.page_width = Cm(29.7), Cm(21)


def field(paragraph, instr: str):
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def add_runs(paragraph, text: str, size=12):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        run = paragraph.add_run(part[2:-2] if bold else part[1:-1] if code else part)
        run.bold = bold
        run.font.name = FONT
        run.font.size = Pt(size)


def body(doc, text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, spacing=1.5, after=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing = spacing
    pf.first_line_indent = Cm(1.25) if indent else Cm(0)
    pf.space_after = Pt(after)
    add_runs(p, text, size)
    return p


def centered(doc, text, bold=False, size=12, before=0, after=0, caps=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text.upper() if caps else text)
    r.bold, r.font.name, r.font.size = bold, FONT, Pt(size)
    return p


def heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text.upper() if level == 1 else text)
    r.bold = True
    r.font.name, r.font.size = FONT, Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def table(doc, rows, number, title):
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    r = cap.add_run(f"Tabela {number} - {title}")
    r.font.name, r.font.size = FONT, Pt(10)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            add_runs(p, cell.strip(), 10)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    src = doc.add_paragraph()
    r = src.add_run("Fonte: elaboração própria, com dados abertos do TSE.")
    r.font.name, r.font.size = FONT, Pt(10)
    src.paragraph_format.space_after = Pt(12)


def parse(md: str):
    lines = md.splitlines()
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c for c in lines[i].strip().strip("|").split("|")]
                if not re.fullmatch(r"[\s:\-|]+", lines[i]):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        if line.startswith("### "):
            blocks.append(("h3", line[4:]))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:]))
        elif line.startswith("# "):
            blocks.append(("title", line[2:]))
        elif re.match(r"^\d+\. ", line):
            blocks.append(("item", line))
        elif line.strip():
            blocks.append(("p", line))
        i += 1
    return blocks


def toc_entries(blocks):
    out = []
    for kind, text in blocks:
        if kind == "h2" and not text.lower().startswith("resumo"):
            out.append((1, "REFERÊNCIAS" if text.lower().startswith("refer") else text.upper()))
        elif kind == "h3":
            out.append((2, text))
    return out


def build(md_path: Path, out: Path, autor: str, instituicao: str, local: str, ano: str,
          pages: dict | None = None):
    blocks = parse(md_path.read_text("utf-8"))
    title = next(t for k, t in blocks if k == "title")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    set_margins(doc.sections[0])

    # Capa
    centered(doc, instituicao, bold=True, caps=True)
    centered(doc, autor, caps=True, before=60)
    titulo, _, subt = title.partition(":")
    centered(doc, titulo.strip(), bold=True, caps=True, before=150)
    if subt:
        centered(doc, subt.strip(), before=6)
    centered(doc, local, before=200)
    centered(doc, ano)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Folha de rosto
    centered(doc, autor, caps=True)
    centered(doc, titulo.strip(), bold=True, caps=True, before=150)
    if subt:
        centered(doc, subt.strip(), before=6)
    nat = doc.add_paragraph()
    nat.paragraph_format.left_indent = Cm(8)
    nat.paragraph_format.space_before = Pt(60)
    nat.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = nat.add_run("Relatório técnico preliminar do Estudo 05 do projeto opianista, sobre o encerramento "
                    "tardio das seções eleitorais no 2º turno das Eleições Gerais de 2026. Versão sujeita a "
                    "revisão independente antes de divulgação.")
    r.font.name, r.font.size = FONT, Pt(10)
    centered(doc, local, before=160)
    centered(doc, ano)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Resumo
    it = iter(blocks)
    for kind, text in it:
        if kind == "h2" and text.lower().startswith("resumo"):
            break
    centered(doc, "RESUMO", bold=True, after=12)
    for kind, text in it:
        if kind == "h2":
            pending = (kind, text)
            break
        if kind == "p":
            body(doc, text, indent=False, spacing=1.5, after=12)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Sumario
    centered(doc, "SUMÁRIO", bold=True, after=12)
    for level, text in toc_entries(blocks):
        tp = doc.add_paragraph()
        tp.paragraph_format.tab_stops.add_tab_stop(Cm(16), alignment=2, leader=WD_TAB_LEADER.DOTS)
        tp.paragraph_format.line_spacing = 1.5
        tp.paragraph_format.space_after = Pt(0)
        r = tp.add_run(f"{text}\t{(pages or {}).get(text, '')}")
        r.font.name, r.font.size, r.bold = FONT, Pt(12), level == 1

    # Texto: nova secao com numeracao no canto superior direito
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(sec)
    sec.header.is_linked_to_previous = False
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field(hp, "PAGE")
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "4")  # folha de rosto = 1, resumo = 2, sumario = 3
    sec._sectPr.append(pg)
    doc.sections[0].header.is_linked_to_previous = False

    tab_n = 0
    last_heading = ""
    stream = [pending] + list(it)
    for kind, text in stream:
        if kind == "h2":
            if text.lower().startswith("refer"):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                p = centered(doc, "REFERÊNCIAS", bold=True, after=12)
                p.style = doc.styles["Heading 1"]
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.font.name, run.font.size, run.bold = FONT, Pt(12), True
                refs = True
                continue
            heading(doc, text, 1)
            last_heading = re.sub(r"^\d+(\.\d+)*\s+", "", text)
        elif kind == "h3":
            heading(doc, text, 2)
            last_heading = re.sub(r"^\d+(\.\d+)*\s+", "", text)
        elif kind == "table":
            tab_n += 1
            table(doc, text, tab_n, last_heading)
        elif kind == "item":
            body(doc, text)
        elif kind == "p":
            if text.startswith("BRASIL.") or text.startswith("**Palavras"):
                body(doc, text, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0, after=12)
            else:
                body(doc, text)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return blocks


def measure_pages(docx_path: Path, blocks) -> dict:
    """Converte para PDF com LibreOffice e localiza a pagina impressa de cada titulo."""
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(docx_path)],
                       check=True, capture_output=True, timeout=240)
        pdf = Path(tmp) / (docx_path.stem + ".pdf")
        n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True,
                                                             text=True).stdout).group(1))
        texts = [subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), "-layout", str(pdf), "-"],
                                capture_output=True, text=True).stdout for i in range(1, n + 1)]
    pages, start = {}, 5  # paginas 1 a 4 fisicas: capa, rosto, resumo, sumario
    for _, text in toc_entries(blocks):
        key = re.sub(r"\s+", " ", text)
        for i in range(start - 1, n):
            if key in re.sub(r"\s+", " ", texts[i]):
                pages[text] = i  # pagina fisica i+1; capa nao conta: numero impresso = i
                start = i + 1
                break
    return pages


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--md", type=Path, required=True)
    p.add_argument("--out", type=Path, required=True)
    p.add_argument("--autor", default="João Pedro Pereira Passos")
    p.add_argument("--instituicao", default="Projeto opianista")
    p.add_argument("--local", default="Palmas")
    p.add_argument("--ano", default="2026")
    a = p.parse_args(argv)
    blocks = build(a.md, a.out, a.autor, a.instituicao, a.local, a.ano)
    try:
        pages = measure_pages(a.out, blocks)
        build(a.md, a.out, a.autor, a.instituicao, a.local, a.ano, pages)
    except (OSError, subprocess.SubprocessError, AttributeError) as e:
        print("Aviso: sumario sem paginas (", e, ")")
    print("OK:", a.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
